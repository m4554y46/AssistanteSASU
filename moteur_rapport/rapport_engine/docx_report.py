from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARKBLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(86, 95, 108)
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(155, 28, 28)
ACCENT = RGBColor(0, 91, 150)
WHITE = RGBColor(255, 255, 255)
PALEBLUE = "E8EEF5"
LIGHT = "F2F4F7"
BORDER_GRAY = "BFBFBF"
ACCENT_BG = "005B96"

_V = {
    "exec_summary": "Cette semaine, j'ai prepare pour toi une selection d'articles et d'opportunites trouves pour ASTRA MOMENTUM.",
    "subtitle": "Veille, prospection et ressources de la semaine",
    "veille_intro": "Voici les articles et ressources retenus cette semaine.",
}

def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url):
    if not url:
        return
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[idx])


def para(doc, text="", style=None, before=None, after=None, color=None, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def first_sentence(text: str, fallback: str) -> str:
    clean = " ".join((text or "").split())
    if not clean:
        return fallback
    clean = clean.replace("\u2014", "-").replace("\u2013", "-")
    clean = re.sub(r"(?i)^listen to this episode on[:]?\s*", "", clean)
    clean = re.sub(r"(?i)^spotify\s*\|.*?(apple\s*)?podcasts?\s*", "", clean)
    clean = re.sub(r"(?i)^in our last article,?\s*", "", clean)
    parts = clean.split(". ")
    return (parts[0].strip() + ".")[:420]


def add_page_number(footer):
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run1 = fp.add_run("ASTRA MOMENTUM | Page ")
    set_run_font(run1, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2 = fp.add_run()
    run2._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    instr.set(qn("xml:space"), "preserve")
    run3 = fp.add_run()
    run3._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run4 = fp.add_run()
    run4._r.append(fld2)


def alt_row_shade(table, row_idx, fill=LIGHT):
    if row_idx % 2 == 1:
        for cell in table.rows[row_idx].cells:
            shade_cell(cell, fill)


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    for edge, attrs in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        if attrs is None:
            continue
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)
        existing = tcBorders.find(qn(f"w:{edge}"))
        if existing is not None:
            tcBorders.remove(existing)
        el = OxmlElement(f"w:{edge}")
        for attr, val in attrs.items():
            el.set(qn(f"w:{attr}"), str(val))
        tcBorders.append(el)


def apply_table_borders(table, sz="4", color=BORDER_GRAY):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"sz": sz, "val": "single", "color": color},
                            bottom={"sz": sz, "val": "single", "color": color},
                            left={"sz": sz, "val": "single", "color": color},
                            right={"sz": sz, "val": "single", "color": color})


def style_header_row(table, headers, fill=PALEBLUE, font_color=NAVY, size=9):
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, fill)
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, bold=True, color=font_color, size=size)
    apply_table_borders(table)


def make_pro_table(doc, headers, rows, widths, header_fill=PALEBLUE, alt_fill=LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    style_header_row(table, headers, fill=header_fill)
    for ri, row_data in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row_data):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    set_run_font(r, size=9)
        alt_row_shade(table, ri + 1, fill=alt_fill)
    apply_table_borders(table)
    if widths:
        set_table_widths(table, widths)
    return table


def accent_heading(doc, text, level=1):
    sizes = {1: (13, 15), 2: (11.5, 12.5), 3: (10.5, 11.5)}
    font_sz = sizes.get(level, (10, 11.5))[1]
    h = doc.add_paragraph(style=f"Heading {level}")
    h.clear()
    r = h.add_run(text)
    set_run_font(r, size=font_sz, color=NAVY, bold=True)
    pPr = h._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "28")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ACCENT_BG)
    pBdr.append(left)
    return h


def add_callout(doc, text, title=None):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "F0F4FA")
    cell.paragraphs[0].clear()
    if title:
        r = cell.paragraphs[0].add_run(title + "\n")
        set_run_font(r, size=10, color=NAVY, bold=True)
    r = cell.paragraphs[0].add_run(text)
    set_run_font(r, size=10, color=RGBColor(51, 51, 51))
    set_cell_margins(cell, top=80, start=120, bottom=80, end=120)
    set_cell_border(cell,
        left={"sz": "28", "val": "single", "color": ACCENT_BG},
        top={"sz": "4", "val": "single", "color": BORDER_GRAY},
        bottom={"sz": "4", "val": "single", "color": BORDER_GRAY},
        right={"sz": "4", "val": "single", "color": BORDER_GRAY})
    para(doc, "", after=6)
    return table


def add_separator(doc, before=6, after=6):
    p = para(doc, after=after, before=before)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BORDER_GRAY)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    doc.core_properties.title = "Rapport Hebdomadaire ASTRA MOMENTUM"
    doc.core_properties.author = "Michael ASSAYAG"
    doc.core_properties.last_modified_by = "Virginie Benayoun"
    doc.core_properties.company = "ASTRA MOMENTUM"
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color, before, after in [
        ("Heading 1", 15, BLUE, 14, 6),
        ("Heading 2", 12.5, BLUE, 10, 4),
        ("Heading 3", 11.5, DARKBLUE, 6, 2),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "ASTRA MOMENTUM | Rapport Hebdomadaire"
    set_run_font(header.runs[0], size=9, color=GRAY)
    hPPr = header._p.get_or_add_pPr()
    hBdr = OxmlElement("w:pBdr")
    hBottom = OxmlElement("w:bottom")
    hBottom.set(qn("w:val"), "single")
    hBottom.set(qn("w:sz"), "4")
    hBottom.set(qn("w:space"), "1")
    hBottom.set(qn("w:color"), BORDER_GRAY)
    hBdr.append(hBottom)
    hPPr.append(hBdr)
    add_page_number(section.footer)
    return doc


def add_dashboard(doc, articles, opportunities, meta):
    accent_heading(doc, "1. Tableau de bord")
    bullet(doc, f"{len(articles)} articles et {len(opportunities)} opportunites retenus cette semaine.")

    opp_label = f"Aucune opportunite" if len(opportunities) == 0 else f"{len(opportunities)} opportunites retenues"
    opp_priority = "N/A" if len(opportunities) == 0 else "Haute"

    make_pro_table(doc,
        ["Signal", "Observation", "Action conseillee", "Priorite"],
        [
            ["Veille IA/Product", f"{len(articles)} contenus retenus", "Enrichir le discours commercial", "Haute"],
            ["Prospection", opp_label, "Qualifier les 2 meilleurs fits", opp_priority],
        ],
        [1.4, 2.2, 2.1, 0.8],
    )


def _norm_title(title: str) -> str:
    return (title or "").replace("\u2014", "-").replace("\u2013", "-")


def add_article(doc, idx, item):
    accent_heading(doc, f"{idx}. {_norm_title(item['title'])}", level=3)
    p = para(doc, after=2)
    set_run_font(p.add_run(f"Source : {item['source']}"), size=9.5, color=GRAY, bold=True)
    p2 = para(doc, after=4)
    add_hyperlink(p2, item["url"], item["url"])
    resume = first_sentence(item.get("snippet", ""), "Aucun extrait disponible.")
    if resume != "Aucun extrait disponible.":
        bullet(doc, resume)


def add_opportunity(doc, idx, item):
    accent_heading(doc, f"{idx}. {_norm_title(item['title'])}", level=3)
    p = para(doc, after=2)
    set_run_font(p.add_run(f"Plateforme : {item['source']}"), size=9.5, color=GRAY, bold=True)
    p2 = para(doc, after=4)
    add_hyperlink(p2, item["url"], item["url"])
    resume = first_sentence(item.get("snippet", ""), "Aucun descriptif disponible.")
    if resume != "Aucun descriptif disponible.":
        bullet(doc, resume)


def add_methods(doc, methods):
    accent_heading(doc, "4. Ressources et methodes")
    for method in methods:
        accent_heading(doc, method["name"], level=3)
        bullet(doc, method["definition"])
        bullet(doc, f"Interet pour les clients : {method['business_value']}")
        bullet(doc, f"Quand l'utiliser : {method['client_use']}")
        bullet(doc, f"Comment en parler : {method['commercial_use']}")
        p = para(doc, after=4)
        set_run_font(p.add_run("Source : "), size=9.5, color=GRAY, bold=True)
        add_hyperlink(p, method["url"], method["url"])


def add_actions(doc, articles, opportunities):
    accent_heading(doc, "5. Priorites de la semaine")
    top_opp = _norm_title(opportunities[0]["title"]) if opportunities else "la meilleure opportunite identifiee"
    top_article = _norm_title(articles[0]["title"]) if articles else "le signal IA/Product le plus fort"
    if opportunities:
        bullet(doc, f"A qualifier en premier : {top_opp}.")
    if articles:
        bullet(doc, f"Article a garder sous le coude pour tes prochains echanges : '{top_article}'.")
    if not opportunities and not articles:
        bullet(doc, "Rien de nouveau a traiter cette semaine. Je continue de surveiller les publications.")


def add_sources(doc, articles, opportunities, methods, meta):
    pass


def _repo_name(item: dict) -> str:
    title = item.get("title", "").replace("\u2014", "-").replace("\u2013", "-")
    for prefix in ["GitHub - ", "GitHub: ", "github.com/"]:
        if title.startswith(prefix):
            title = title[len(prefix):]
    parts = title.split(": ")
    return parts[0].strip() if len(parts) > 1 else title[:60].strip()


def add_tokenforge_watch(doc, items: list[dict]):
    if not items:
        return
    add_separator(doc)
    accent_heading(doc, "7. TokenForge Watch - Revue de presse des couts IA")
    para(doc,
         "Actualites, outils et mouvements de marche sur la reduction des factures IA en entreprise.",
         before=4, after=8, color=GRAY, size=9.5, italic=True)

    github_repos = [i for i in items if i.get("kind") == "github" or "github.com" in i.get("url", "")]
    articles_news = [i for i in items if i not in github_repos]

    if articles_news:
        accent_heading(doc, "A. Marche et actualites des prix", level=2)
        for a in articles_news[:5]:
            title = a.get("title", "")[:120].replace("\u2014", "-").replace("\u2013", "-")
            resume = first_sentence(a.get("snippet", ""), "")
            p = para(doc, after=2)
            set_run_font(p.add_run(title), size=10, bold=True, color=NAVY)
            p2 = para(doc, after=2)
            if resume:
                set_run_font(p2.add_run(resume + " "), size=9, color=GRAY)
            if a.get("url"):
                add_hyperlink(p2, a["url"], a["url"])

    if github_repos:
        accent_heading(doc, "B. Outils et initiatives open source", level=2)
        for g in github_repos[:6]:
            name = _repo_name(g)
            resume = first_sentence(g.get("snippet", ""), "")
            p = para(doc, after=2)
            set_run_font(p.add_run(name + " "), size=10, bold=True, color=NAVY)
            if resume:
                set_run_font(p.add_run(resume + " "), size=9, color=GRAY)
            add_hyperlink(p, g.get("url", ""), g.get("url", ""))

    accent_heading(doc, "C. Ce que j'en retiens pour TokenForge", level=2)
    vendors_news = [i for i in articles_news if any(v in (i.get("title","") + i.get("snippet","")).lower() for v in ["openai", "anthropic", "google", "mistral", "pricing", "tarif", "price"])]
    if vendors_news:
        bullet(doc, "Les fournisseurs ajustent leurs grilles de prix. A suivre pour ajuster les recommandations TokenForge.")
    routers = [g for g in github_repos if any(r in (_repo_name(g)+g.get("snippet","")).lower() for r in ["router", "gateway", "litellm", "proxy", "route"])]
    if routers:
        bullet(doc, f"{len(routers)} outil(s) de routage entre fournisseurs identifie(s) - a regarder pour basculer sur le provider le moins cher en temps reel.")
    compressors = [g for g in github_repos if any(c in (_repo_name(g)+g.get("snippet","")).lower() for c in ["compress", "token", "rtk", "cache"])]
    if compressors:
        bullet(doc, f"{len(compressors)} outil(s) de compression ou de cache identifie(s) - a tester pour reduire les tokens envoyes a l'API.")
    budget = [i for i in items if any(b in (i.get("title","")+i.get("snippet","")).lower() for b in ["finops", "budget", "cost", "facture"])]
    if budget:
        bullet(doc, f"{len(budget)} source(s) sur la maitrise des couts IA - utile pour le volet conseil de TokenForge (tableau de bord couts, alertes, plafonds).")
    bullet(doc, "Prochaine etape : monter un petit comparatif des fournisseurs (OpenAI vs Anthropic vs Mistral vs Google) avec les prix token 2026 pour les clients TokenForge.")


def build_report(config: dict, articles: list[dict], opportunities: list[dict], methods: list[dict], meta: dict, output_path: Path, tokenforge_items: list[dict] | None = None) -> Path:
    doc = setup_doc()
    today = date.today().isoformat()
    para(doc, "RAPPORT HEBDOMADAIRE ", size=23, color=NAVY, bold=True, after=4)
    para(doc, _V["subtitle"], size=13.5, color=GRAY, after=14)
    for label, value in [
        ("Destinataire", f"{config['recipient']} - President SASU"),
        ("Autrice", config["author"]),
        ("Positionnement", config["positioning"]),
        ("Date du rapport", today),
    ]:
        p = para(doc, after=2)
        set_run_font(p.add_run(label + " : "), bold=True, color=BLACK)
        set_run_font(p.add_run(value), color=BLACK)

    add_separator(doc, before=10, after=10)

    add_callout(doc, _V["exec_summary"], title="Lecture executive")

    add_dashboard(doc, articles, opportunities, meta)
    add_separator(doc)
    accent_heading(doc, "2. Veille de la semaine")
    para(doc,
         _V["veille_intro"],
         before=4, after=8, color=GRAY, size=9.5, italic=True)

    for idx, item in enumerate(articles, 1):
        add_article(doc, idx, item)
    add_separator(doc)
    accent_heading(doc, "3. Missions et opportunites")
    if not opportunities:
        bullet(doc, "Aucune opportunite identifiee cette semaine.")
    else:
        for idx, item in enumerate(opportunities, 1):
            add_opportunity(doc, idx, item)
    add_methods(doc, methods)
    add_actions(doc, articles, opportunities)
    add_sources(doc, articles, opportunities, methods, meta)
    if tokenforge_items:
        add_tokenforge_watch(doc, tokenforge_items)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
