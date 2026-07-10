from __future__ import annotations

import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from core.common import (
    bullet,
    ensure_output_dir,
    load_config,
    para,
    save_json_report,
    set_run_font,
    set_table_widths,
    shade_cell,
    setup_doc,
)
from docx.shared import Pt, RGBColor

NAVY = RGBColor(11, 37, 69)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(86, 95, 108)
PALEBLUE = "E8EEF5"
LIGHT = "F2F4F7"


MONTHS_FR = [
    "", "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre"
]

INPUT_DIR = Path(__file__).resolve().parent / "notes"


def list_notes() -> list[Path]:
    if not INPUT_DIR.exists():
        INPUT_DIR.mkdir(parents=True, exist_ok=True)
        return []
    files = sorted(INPUT_DIR.glob("*.txt")) + sorted(INPUT_DIR.glob("*.md"))
    return files


def parse_note_file(path: Path) -> dict:
    raw = path.read_text(encoding="utf-8")
    lines = raw.strip().split("\n")

    # Extract metadata from first lines if present
    meta = {
        "file": path.name,
        "client": "",
        "projet": "",
        "date": date.today(),
        "participants": "",
        "duree": "",
    }

    text_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped.lower().startswith("client:"):
            meta["client"] = stripped.split(":", 1)[1].strip()
        elif stripped.lower().startswith("projet:"):
            meta["projet"] = stripped.split(":", 1)[1].strip()
        elif stripped.lower().startswith("date:"):
            d = stripped.split(":", 1)[1].strip()
            for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
                try:
                    meta["date"] = datetime.strptime(d, fmt).date()
                    break
                except ValueError:
                    pass
        elif stripped.lower().startswith("participants:"):
            meta["participants"] = stripped.split(":", 1)[1].strip()
        elif stripped.lower().startswith("durée:") or stripped.lower().startswith("duree:"):
            meta["duree"] = stripped.split(":", 1)[1].strip()
        else:
            text_lines.append(line)

    text = "\n".join(text_lines).strip()
    meta["raw_text"] = text
    return meta


def structure_cr(meta: dict) -> dict:
    text = meta["raw_text"]

    # Extract sections
    sections = {
        "contexte": "",
        "points_abordes": [],
        "decisions": [],
        "actions": [],
        "prochaines_etapes": "",
    }

    # Try to split by section headers
    current_section = "points_abordes"
    for line in text.split("\n"):
        l = line.strip().lower()
        if re.match(r"^(contexte|objectif|contexte de la réunion)", l):
            current_section = "contexte"
            continue
        elif re.match(r"^(points? (abordés|discutés)|ordre du jour|échanges|discussion)", l):
            current_section = "points_abordes"
            continue
        elif re.match(r"^(décisions?|decisions?|validations?)", l):
            current_section = "decisions"
            continue
        elif re.match(r"^(actions?|todo|à faire|prochaines étapes?|next steps|follow.up)", l):
            current_section = "actions"
            continue
        elif re.match(r"^(prochaine réunion|prochain call|conclusion)", l):
            current_section = "prochaines_etapes"
            continue

        line_clean = line.strip()
        if not line_clean or line_clean.startswith("#") or line_clean.startswith("---"):
            continue

        if line_clean.startswith("- ") or line_clean.startswith("* "):
            content = line_clean[2:]
            if current_section in ("points_abordes",):
                sections["points_abordes"].append(content)
            elif current_section in ("decisions",):
                sections["decisions"].append(content)
            elif current_section in ("actions",):
                sections["actions"].append(content)
        elif current_section == "contexte":
            sections["contexte"] += line_clean + " "
        elif current_section == "prochaines_etapes":
            sections["prochaines_etapes"] += line_clean + " "
        else:
            sections["points_abordes"].append(line_clean)

    sections["contexte"] = sections["contexte"].strip()
    sections["prochaines_etapes"] = sections["prochaines_etapes"].strip()

    # If no structured sections detected, treat everything as points abordes
    if not sections["points_abordes"] and not sections["decisions"] and not sections["actions"]:
        for line in text.split("\n"):
            l = line.strip()
            if l and not l.startswith("#") and not l.startswith("---"):
                sections["points_abordes"].append(l)

    # Parse actions for owner/deadline
    parsed_actions = []
    for action in sections["actions"]:
        owner = extract_owner(action)
        deadline = extract_deadline(action)
        parsed_actions.append({
            "description": action,
            "owner": owner,
            "deadline": deadline,
        })
    sections["actions_parsed"] = parsed_actions

    return sections


def extract_owner(text: str) -> str:
    patterns = [
        r"\[([^\]]+)\]",
        r"\(([^)]+)\)",
        r"(?:pour|par|responsable|owner)\s*:?\s*([A-Z][a-zéèêëàâîïôûù]+(?:\s+[A-Z][a-zéèêëàâîïôûù]+)?)",
    ]
    for pat in patterns:
        m = re.search(pat, text)
        if m:
            name = m.group(1).strip()
            if not any(kw in name.lower() for kw in ["pour", "par", "owner", "faire", "avant", "le"]):
                return name
    return "À définir"


def extract_deadline(text: str) -> str:
    patterns = [
        r"(\d{2}/\d{2}/\d{4})",
        r"(\d{4}-\d{2}-\d{2})",
        r"(avant\s+(?:le\s+)?(\d{1,2}\s+[a-zéèêëàâîïôûù]+\s+\d{4}))",
        r"(semaine\s+prochaine)",
        r"(cette\s+semaine)",
        r"(courant\s+[a-zéèêëàâîïôûù]+)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            return m.group(0)
    return "Non spécifiée"


def build_report(config: dict, meta: dict, sections: dict, output_path: Path):
    client = meta.get("client", "Client")
    projet = meta.get("projet", "Réunion")
    d = meta.get("date", date.today())

    doc = setup_doc(
        f"COMPTE-RENDU DE RÉUNION",
        f"{projet} - {client}",
        config.get("author", "Assistant IA SASU"),
    )

    meta_table = doc.add_table(rows=1, cols=2)
    set_table_widths(meta_table, [1.5, 4.5])
    meta_data = [
        ("Client", client),
        ("Projet", projet),
        ("Date", d.isoformat()),
        ("Participants", meta.get("participants", "Non renseigné")),
        ("Durée", meta.get("durée", meta.get("duree", "Non renseignée"))),
        ("Rédigé le", date.today().isoformat()),
    ]
    for i, (label, value) in enumerate(meta_data):
        shade_cell(meta_table.rows[0].cells[i % 2], LIGHT)
        if i == 0:
            meta_table.rows[0].cells[0].text = label
            meta_table.rows[0].cells[1].text = value
        else:
            cells = meta_table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        for row in meta_table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r, size=10)

    para(doc, "", after=8)

    # Contexte
    if sections.get("contexte"):
        para(doc, "1. Contexte et objectif", style="Heading 1")
        para(doc, sections["contexte"])

    # Points abordés
    para(doc, "2. Points abordés", style="Heading 1")
    for idx, point in enumerate(sections.get("points_abordes", []), 1):
        bullet(doc, point)

    # Décisions
    if sections.get("decisions"):
        para(doc, "3. Décisions", style="Heading 1")
        for dec in sections["decisions"]:
            bullet(doc, dec)

    # Actions
    actions = sections.get("actions_parsed", [])
    if actions:
        para(doc, "4. Actions et responsabilités", style="Heading 1")
        action_table = doc.add_table(rows=1, cols=3)
        for i, h in enumerate(["Action", "Responsable", "Échéance"]):
            cell = action_table.rows[0].cells[i]
            shade_cell(cell, PALEBLUE)
            cell.text = h
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, bold=True, color=NAVY, size=9.5)
        for a in actions:
            cells = action_table.add_row().cells
            cells[0].text = a["description"]
            cells[1].text = a["owner"]
            cells[2].text = a["deadline"]
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r, size=9.2)
        set_table_widths(action_table, [3.5, 1.5, 1.5])
    elif sections.get("actions"):
        para(doc, "4. Actions", style="Heading 1")
        for a in sections["actions"]:
            bullet(doc, a)

    if sections.get("prochaines_etapes"):
        para(doc, "5. Prochaines étapes", style="Heading 1")
        para(doc, sections["prochaines_etapes"])

    doc.save(output_path)
    return output_path


def main() -> int:
    config = load_config()
    agent_cfg = config.get("cr_reunion", {})
    if not agent_cfg.get("enabled", True):
        print("Agent CR Réunion désactivé.")
        return 0

    print(">>> Agent 4 : Compte-rendu de réunion")
    notes = list_notes()

    if not notes:
        print("  Aucune note trouvée.")
        print(f"  Déposez un fichier .txt dans : {INPUT_DIR}")
        print("  Exemple de format :")
        print("    Client: ACME Corp")
        print("    Projet: Audit IA")
        print("    Date: 2026-06-03")
        print("    Participants: Jean, Marie")
        print("    ---")
        print("    Points abordés:")
        print("    - Point 1...")
        print("    Décisions:")
        print("    - Décision 1...")
        print("    Actions:")
        print("    - [Michel] Faire X avant 15/06/2026")
        return 0

    output_dir = ensure_output_dir(config)
    today = date.today().isoformat()
    generated = []

    for note_path in notes:
        meta = parse_note_file(note_path)
        sections = structure_cr(meta)
        client_slug = meta.get("client", "").replace(" ", "_") or "reunion"
        output_path = output_dir / f"CR_{client_slug}_{today}.docx"
        build_report(config, meta, sections, output_path)
        generated.append(str(output_path))
        print(f"  OK {note_path.name} -> {output_path.name}")
        # Archive processed note
        archived = note_path.parent / "processed"
        archived.mkdir(exist_ok=True)
        note_path.rename(archived / note_path.name)

    run_log = output_dir / f"CR_Reunion_{today}_log.json"
    save_json_report(
        {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "notes_traitees": len(notes),
            "outputs": generated,
        },
        run_log,
    )

    print(f"  {len(generated)} CR généré(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
