from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARKBLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(86, 95, 108)
BLACK = RGBColor(0, 0, 0)
ACCENT = RGBColor(0, 91, 150)
WHITE = RGBColor(255, 255, 255)
PALEBLUE = "E8EEF5"
LIGHT = "F2F4F7"
BORDER_GRAY = "BFBFBF"
ACCENT_BG = "005B96"


ROOT = Path(__file__).resolve().parent.parent
CONFIG_PATH = ROOT / "config.json"


def load_config() -> dict:
    with CONFIG_PATH.open("r", encoding="utf-8") as f:
        return json.load(f)


def load_agent_config(agent_name: str) -> dict:
    cfg = load_config()
    return cfg.get(agent_name, {})


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url):
    if not url:
        return
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Pt(width * 12)
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[idx])


def para(doc, text="", style=None, before=None, after=None, color=None, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_page_number(footer):
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run1 = fp.add_run("ASTRA MOMENTUM | Page ")
    set_run_font(run1, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2 = fp.add_run()
    run2._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    instr.set(qn("xml:space"), "preserve")
    run3 = fp.add_run()
    run3._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run4 = fp.add_run()
    run4._r.append(fld2)


def set_doc_meta(doc, title, creator, last_modified_by, company):
    core_props = doc.core_properties
    core_props.title = title
    core_props.author = creator
    core_props.last_modified_by = last_modified_by
    core_props.company = company


def alt_row_shade(table, row_idx, fill=LIGHT):
    if row_idx % 2 == 1:
        for cell in table.rows[row_idx].cells:
            shade_cell(cell, fill)


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    for edge, attrs in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        if attrs is None:
            continue
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)
        existing = tcBorders.find(qn(f"w:{edge}"))
        if existing is not None:
            tcBorders.remove(existing)
        el = OxmlElement(f"w:{edge}")
        for attr, val in attrs.items():
            el.set(qn(f"w:{attr}"), str(val))
        tcBorders.append(el)


def apply_table_borders(table, sz="4", color=BORDER_GRAY):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"sz": sz, "val": "single", "color": color},
                            bottom={"sz": sz, "val": "single", "color": color},
                            left={"sz": sz, "val": "single", "color": color},
                            right={"sz": sz, "val": "single", "color": color})


def style_header_row(table, headers, fill=PALEBLUE, font_color=NAVY, size=9):
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, fill)
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, bold=True, color=font_color, size=size)
    apply_table_borders(table)


def make_pro_table(doc, headers, rows, widths, header_fill=PALEBLUE, alt_fill=LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    style_header_row(table, headers, fill=header_fill)
    for ri, row_data in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row_data):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    set_run_font(r, size=9)
        alt_row_shade(table, ri + 1, fill=alt_fill)
    apply_table_borders(table)
    if widths:
        set_table_widths(table, widths)
    return table


def accent_heading(doc, text, level=1):
    sizes = {1: (13, 15), 2: (11.5, 12.5), 3: (10.5, 11.5)}
    accent_sz, font_sz = sizes.get(level, (10, 11.5))
    h = doc.add_paragraph(style=f"Heading {level}")
    h.clear()
    r = h.add_run(text)
    set_run_font(r, size=font_sz, color=NAVY, bold=True)
    pPr = h._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "28")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ACCENT_BG)
    pBdr.append(left)
    return h


def add_callout(doc, text, title=None):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "F0F4FA")
    cell.paragraphs[0].clear()
    if title:
        r = cell.paragraphs[0].add_run(title + "\n")
        set_run_font(r, size=10, color=NAVY, bold=True)
    r = cell.paragraphs[0].add_run(text)
    set_run_font(r, size=10, color=RGBColor(51, 51, 51))
    set_cell_margins(cell, top=80, start=120, bottom=80, end=120)
    set_cell_border(cell,
        left={"sz": "28", "val": "single", "color": ACCENT_BG},
        top={"sz": "4", "val": "single", "color": BORDER_GRAY},
        bottom={"sz": "4", "val": "single", "color": BORDER_GRAY},
        right={"sz": "4", "val": "single", "color": BORDER_GRAY})
    para(doc, "", after=6)
    return table


def add_separator(doc, before=6, after=6):
    p = para(doc, after=after, before=before)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BORDER_GRAY)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def setup_doc(title: str, subtitle: str, author: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(612)
    section.page_height = Pt(792)
    section.top_margin = Pt(57)
    section.bottom_margin = Pt(57)
    section.left_margin = Pt(61)
    section.right_margin = Pt(61)
    header = section.header.paragraphs[0]
    header.text = "ASTRA MOMENTUM"
    set_run_font(header.runs[0], size=9, color=GRAY, italic=True)
    hPPr = header._p.get_or_add_pPr()
    hBdr = OxmlElement("w:pBdr")
    hBottom = OxmlElement("w:bottom")
    hBottom.set(qn("w:val"), "single")
    hBottom.set(qn("w:sz"), "4")
    hBottom.set(qn("w:space"), "1")
    hBottom.set(qn("w:color"), BORDER_GRAY)
    hBdr.append(hBottom)
    hPPr.append(hBdr)
    add_page_number(section.footer)
    set_doc_meta(doc, title, "Michael ASSAYAG", author, "ASTRA MOMENTUM")

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color, before, after in [
        ("Heading 1", 15, BLUE, 14, 6),
        ("Heading 2", 12.5, BLUE, 10, 4),
        ("Heading 3", 11.5, DARKBLUE, 6, 2),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    para(doc, title, size=22, color=NAVY, bold=True, after=4)
    para(doc, subtitle, size=13, color=GRAY, after=14)
    p = para(doc, after=2)
    set_run_font(p.add_run("Autrice : "), bold=True, color=BLACK)
    set_run_font(p.add_run(author), color=BLACK)
    p = para(doc, after=2)
    set_run_font(p.add_run("Date : "), bold=True, color=BLACK)
    set_run_font(p.add_run(date.today().isoformat()), color=BLACK)
    add_separator(doc, before=10, after=10)

    return doc


def ensure_output_dir(cfg: dict, key: str = "output_dir") -> Path:
    p = Path(cfg.get(key, ROOT / "output"))
    p.mkdir(parents=True, exist_ok=True)
    return p


def save_json_report(data: Any, path: Path):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
