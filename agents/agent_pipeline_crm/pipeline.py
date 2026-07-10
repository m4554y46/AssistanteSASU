from __future__ import annotations

import csv
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from core.common import (
    add_hyperlink,
    bullet,
    ensure_output_dir,
    load_config,
    para,
    save_json_report,
    set_run_font,
    set_table_widths,
    shade_cell,
    setup_doc,
)
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(86, 95, 108)
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(155, 28, 28)
GREEN = RGBColor(0, 128, 0)
ORANGE = RGBColor(200, 120, 0)
PALEBLUE = "E8EEF5"
LIGHTGRAY = "F2F4F7"


def load_pipeline(csv_path: str) -> list[dict]:
    path = Path(csv_path)
    rows = []
    if not path.exists():
        return rows
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            rows.append(row)
    return rows


def days_since(date_str: str | None) -> int | None:
    if not date_str:
        return None
    try:
        d = datetime.strptime(date_str.strip(), "%Y-%m-%d").date()
        return (date.today() - d).days
    except ValueError:
        return None


def stage_index(stage: str, stages: list[str]) -> int:
    for i, s in enumerate(stages):
        if s.lower() == stage.lower().strip():
            return i
    return -1


def status_color(days: int | None) -> tuple[str, RGBColor]:
    if days is None:
        return "À relancer", ORANGE
    if days >= 14:
        return "URGENT - Relance nécessaire", RED
    if days >= 7:
        return "À surveiller", ORANGE
    return "OK - Récent", GREEN


def build_report(config: dict, pipeline: list[dict], stages: list[str], output_path: Path):
    doc = setup_doc(
        "SUIVI PIPELINE CRM - RAPPORT HEBDOMADAIRE",
        "État du pipeline commercial, relances et actions prioritaires",
        config.get("author", "Assistant IA SASU"),
    )

    para(doc, "1. Résumé du pipeline", style="Heading 1")
    total = len(pipeline)
    if total == 0:
        bullet(doc, "Aucun prospect dans le pipeline. Ajoutez des lignes dans le fichier pipeline.csv.")
        doc.save(output_path)
        return

    # Stats
    stages_count = {}
    for s in stages:
        stages_count[s] = 0
    for p in pipeline:
        st = p.get("stage", "").strip()
        found = False
        for s in stages:
            if st.lower() == s.lower():
                stages_count[s] += 1
                found = True
                break
        if not found:
            stages_count[st] = stages_count.get(st, 0) + 1

    para(doc, f"Total prospects dans le pipeline : {total}", size=12, bold=True, color=NAVY)
    for s, count in stages_count.items():
        if count > 0:
            bullet(doc, f"{s} : {count} prospect(s)")

    # Relances urgentes
    urgent = [p for p in pipeline if days_since(p.get("date_dernier_contact")) is not None and days_since(p.get("date_dernier_contact")) >= 7]
    if urgent:
        para(doc, f"Relances urgentes cette semaine : {len(urgent)}", bold=True, color=RED, size=11, before=10)

    para(doc, "", after=8)

    # Section 2: Tableau détaillé
    para(doc, "2. Détail du pipeline", style="Heading 1")
    table = doc.add_table(rows=1, cols=6)
    headers = ["Entreprise", "Contact", "Stage", "Dernier contact", "Statut", "Prochaine action"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, PALEBLUE)
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, bold=True, color=NAVY, size=8.5)

    for row_data in pipeline:
        cells = table.add_row().cells
        ds = days_since(row_data.get("date_dernier_contact"))
        status, color = status_color(ds)
        vals = [
            row_data.get("entreprise", ""),
            row_data.get("contact", ""),
            row_data.get("stage", ""),
            row_data.get("date_dernier_contact", ""),
            status,
            row_data.get("prochaine_action", ""),
        ]
        for i, v in enumerate(vals):
            cells[i].text = v
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_run_font(r, size=8.5, color=RED if i == 4 and ds and ds >= 14 else BLACK)
        if ds and ds >= 14:
            for cell in cells:
                shade_cell(cell, "FDE8E8")

    set_table_widths(table, [1.2, 1.1, 1.2, 0.9, 0.9, 1.2])

    # Section 3: Recommandations de relance
    para(doc, "3. Recommandations de relance", style="Heading 1")
    for p in pipeline:
        ds = days_since(p.get("date_dernier_contact"))
        if ds is None or ds < 7:
            continue
        ent = p.get("entreprise", "Inconnu")
        contact = p.get("contact", "Contact")
        notes = p.get("notes", "")
        para(doc, f"{ent} - {contact}", style="Heading 3")
        bullet(doc, f"Dernier contact il y a {ds} jours")
        if notes:
            bullet(doc, f"Contexte : {notes[:200]}")
        bullet(doc, f"Action recommandée : {suggerer_relance(ds, p)}")

    doc.save(output_path)
    return output_path


def suggerer_relance(days: int, p: dict) -> str:
    stage = p.get("stage", "").lower()
    if "premier contact" in stage:
        return "Relancer avec un cas client concret ou un article de veille pertinent. Proposer un call de 15 min."
    if "rendez-vous" in stage:
        return "Envoyer la proposition commerciale si ce n'est pas fait. Relancer avec un résumé des points clés du call."
    if "proposition" in stage:
        return "Demander un retour sur la proposition. Proposer un point téléphonique pour répondre aux questions."
    if "négociation" in stage:
        return "Relancer courtoisement sur les prochaines étapes. Proposer une date de démarrage."
    return "Relance standard : prendre des nouvelles, partager un contenu à valeur ajoutée."


def main() -> int:
    config = load_config()
    agent_cfg = config.get("pipeline_crm", {})
    if not agent_cfg.get("enabled", True):
        print("Agent Pipeline CRM désactivé.")
        return 0

    print(">>> Agent 2 : Suivi Pipeline CRM")
    csv_path = agent_cfg.get("csv_path", "")
    stages = agent_cfg.get("stages", [])
    pipeline = load_pipeline(csv_path)
    print(f"  Prospects chargés: {len(pipeline)}")

    output_dir = ensure_output_dir(config)
    today = date.today().isoformat()
    output_path = output_dir / f"Pipeline_CRM_{today}.docx"
    build_report(config, pipeline, stages, output_path)

    run_log = output_dir / f"Pipeline_CRM_{today}_log.json"
    save_json_report(
        {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "prospects_count": len(pipeline),
            "pipeline": pipeline,
        },
        run_log,
    )

    print(f"  Rapport généré : {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
